import streamlit as st
import requests
import base64
import io
import csv
import html
import hashlib
import re
import time
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
from fpdf import FPDF

# ---------------------------------------------------------
# Asistente de Escritura Académica
# Proyecto de investigación liderado por la Universidad Nacional
# de Educación (UNAE), en piloto con la pasarela RAG/LLM de CEDIA
# (colección "escritura_academica")
# ---------------------------------------------------------

CEDIA_BASE_URL = "https://ai.hpc.cedia.edu.ec"
COLLECTION = "escritura_academica"

# Logo institucional de la UNAE, en base64 para que la app no dependa
# de un archivo de imagen aparte en el repositorio.
LOGO_UNAE_B64 = "iVBORw0KGgoAAAANSUhEUgAAAiIAAADPCAYAAADI8bpfAAAABHNCSVQICAgIfAhkiAAAAAFzUkdCAK7OHOkAAAAEZ0FNQQAAsY8L/GEFAAAACXBIWXMAAC4jAAAuIwF4pT92AAAAGXRFWHRTb2Z0d2FyZQB3d3cuaW5rc2NhcGUub3Jnm+48GgAAACF0RVh0Q3JlYXRpb24gVGltZQAyMDIyOjEyOjEyIDE2OjUxOjU35qLocQAAUCtJREFUeF7tnQmYXEXV98+5PZONTUAWEVFRURBETWTLLIAJhMxMECECsruh4sL3KiIquIMKLqio6OuC7ATBZHpCWONMT8IiQVFAVBR42YIssmSbpe/5/tV9ZjLdt25v0zPdkzm/57l961R336Vu3apTVadOkWEYhmEYhmEYhmEYxqSDdZ9PXLxhTDZE94ZhGMYYEFE42rnlz9jtnZUMY3IjQrt1Uc8/VTQMwzCqTKB7wzAMwzCMcadIj4jc5z6yYcOYLPBr8bGlC1mPiGEYxthSUBF5WXh6N3VvcGHDmCy0cfMSJu5wYVNEDMMwxhYbmjEMwzAMo2aYImIYhmEYRs0wRcQwDMMwjJphiohhGIZhGDXDFBHDMAzDMGqGKSKGYRiGYdQMU0QMwzAMw6gZpogYhmEYhlEzTBExDMMwDKNmmCJiGIZhGEbNMEXEMAzDMIyaYYqIYRiGYRg1wxQRwzAMwzBqhikihmEYhmHUDFNEDMMwDMOoGaaIGIZhGIZRM0wRMQzDMAyjZpgiYhiGYRhGzTBFxDAMwzCMmmGKiGEYhmEYNcMUEcMwDMMwaoYpIoZhGIZh1AxTRAzDMAzDqBmmiBiGYRiGUTNMETEMwzAMo2aYImIYhmEYRs0wRcQwDMMwjJphiohhGIZhGDXDFBHDMAzDMGoG636Ydm75M3Z7u/DLwtO7qXuDC08m2qhp6zQ1bqPiJCXsW0bdj6swqWjj5iVM3OHCIrRbF/X8M/NFFeig5teHJJH3zodQYt1S6l6tYoR5tM+WCZryShULkqBwwxJa+aSKGebTQa9lGkiomENALJ2UeljFkminmTOEpu+oYkECmrK2k257WsUM86l1R1zM5ioWZIACWUbL/6UiLaQ9pqyjbXZWsSRmUOI/i6h7jYo5oAzYVYNFaaBEuJi6H1GRjqL9pq+nhlepOBb8t4t6/6vhYdpp9i5C3KBiWTAlXppOz7y0iB7o16iq4dKjj6Y0CcluQrQN3q0XEf3wdOLlI9N/AR2wU5qCaSpSmvqfXUZ3vaRihuyxpr5axaKsp8ZnbqFb3PmGKecddLjnO4Xo2bi8Ui0Oof22aaSGV6hYNrOo95GvouBWsSCVnMu9s++k254p9RzlYIqIh3ZqOQcpg/SevKDAuL9LevZUcVIxlopIG7f04aVDuVYcFNw3dElqvooRkE8/jnx6kYrF6ElKT6uGM+Bdd4pJbIUpwh/vou6fqlgUXE8HrmeJikWQ65OSeq8KGdq5eTGKpAUqFgX3M1x+tdHsvZkTruwqGRE5votSl6uYA9LGFbalVlbP41q21TCupfVAZlmuYtVBvvga8sWXVRwGeetRXPAuKlaIuMbHPcj3ywNqvDJfWSwHp+gGPPhFKEfH47qma/QIZAO+uyyQhi+58+D5Q5nj1+qX7gF9NEmpi1XK0EEtc6FC3KRicUQ+gWPkvCNIp7W4nhkqloG4uvBZXPM/EL4dabT0XZS6o1oVM+7/W7j/M1UsG5FwG5+C6gP5+1zszspKZYGqwaUBPYXgKpZgeT/1dd1Edzyf/boybGjGMAwvzOH3D6PWWSoakwLeGdsCZv6+8MDjqLB+cwTtv71+WTLt1Hw08+B9ONaH/UqIg6fhuw+FPPgXp2BA3qiE1CU8DdvOuOaD0VD5YsC84m5u/td8ajlzIbWW1JO3CYDbp+3w8TakwSnE8ttGnvIU8smVeObvzP6kfEwRMQwjBp4asPzuPfTu4da+MZnIDPOcNMCND6KyPSgbV5x2aj0BWuyVqKhKqpxRqW0vLEtVnFDgHl8XMH1rHYcPZe57EoLn53p4j8EzvxsKyaVIh5KGi0diiohhGLGgkNllkAcuyQaNScrWzJR0Q04qx9JBs98KpeJ/ESwzv1Rm21IvQCHZwfUOtHHLVc5WSqMnG+6ZH08c/nU+NR+QjSoNU0QMYxxhkS+GQp8nklhbCiH6fuY3Evxao7zg5b3d/S57PPKO0eJYF2WPRTlj7RlEvqn/vSobEUtbOzUXHbtGYfy3jdcjyzQ6B1zP+e57Eb5UozYi8mtsZyDwbY2JICSd7jci9FmNyjCFwqc2nptuzMb6kOuGfpcg+ZNGRsD3Zw39Dtf8d43OAdfyk8z3Ql/TqAxMg//e+F95UKNzwDGdXYEbz4/d8N8XsC8d4W/pOa/QmAj4zhlc+s6Vxj4W1DAzkHkvmUtzN9MoL8KJb+O3cTZQOA8tdnkS2++xPabxJZOmxD8yeSSTT+gf2dgI/x36TUjBCo0bhoXO1uf6e43y4YxkNX1kbSamBHDvRwvPuNUZkmtUyYgEyzbmG0pqdDHOG/rPGkqs17jiSHjjxnP5nz2uwRkuj8gjpdrC8I4B020d1NSuEUWJaK3tZqxqxqoAmc6MVUG1jVWHQB77FPLYhSrmwJLes5NW3K9iSeC9dTNIIjM9Qgn3W0q9d6ropY2aj2Pmy1T0kimsJJjXRd23aFRBoLh8npjPU3EYFn5TJ3U/pKKXw6hptwQH/sofCgiex3dV9HIYteyVYPqLinlEjWQLcTi1vmKQ5YlMRTwCvB+PrRHeFeXjoEZ5wXNxFUpbVhqJPI7reI0KXhbSwsR6Xv0cUm0rjcqAZ+E1Vh3C2Wcg7b3KJRTAub5nOJNmNu5Em88Ulk9APC4bG0VEvtpFqa+omMOh1PyqBubHkVaeBq78fINM/dzIGSy4gWAVNX8E1/pDiI3Z2BF4jFVH0sYtN+Jch6g4AnkUafs6FWIpZOwtkn57F624V0VqpdZpryDacYDSTQEHh+L5vw/nLmB0LstfluCQYvkjDtzbLTj+u1WMBQ2bD3VS6pcqVgTO5TWgRz77NfLZB1R0cDu1bhtSuAfKi7lMcjSi3qTfeZANUHTevZRSKzUiFusRMQyjKFDMEsThZW6KpUbVLTdQz19RiN6uYh7cPo+atlOhKIMkJ6GQjna1C/2s0kqmVBbRorQQF1Qiq8UqWjUABfGOpPQc71rJGh2Bmd6vwQiNFMxBWkXqFFTaD0yXHT+eP43WzTaBovEzKJdO+alrXIPcTdGGUn8Z0ugEkYSbiny1fu2BD9qCw2+oUBZQfp0SlWOTg3OtxxbpkRDmEzU4HkiSup+FYtGDRurZULT2gMb2YbxrMT13PA0Ky6JSbMxMETEMoyTcOHiaG65A67Dux/NZgl9oMJ9GaFSxLf4ITB/W0AikL6CGUbVCSwUt3pQGx413Uc/5qPRieuT4TfOpZXcVckCFFNfDs9wpVRqOMIN2QFpKrL+cemQpLX8UlfExInKORvn4zHxqfbuGS2aQwpPzFTrIt2H3t6yUQ3M5Pm+qiVPEoUj+L0nQjPziHWbDde80wAOxQ61DmCJiGEbJoGBp3Zzl6yrWLdOo76q4lhpa9SO7m2OZT80tuN+3qjgMCt1rRuNfoxwC4i6ccdHILRAua9iuXNQvxvVZKQrSxN8dzxJjpCkFe6AyPT/CX8Tvfj5yw0X4Kt66ootSX0d+iFFKuYHLf1cYRGffiNyAz96skAOyc3C8hmtCF3XfB4X5PUiHGGd4cvICOuDNKngxRcQwjDie0X0OqIjOzDgvq2OupTuc4V6M0SbvVUpLFSX8qRrMAa3gUp3IjZol1POnpKTeN3LrpJ5r9OuxQ6iQXZTfe66wM2qMgPxyaLFWexf1/Ar3durIzQ0B6Nd1zQzh01EJ53gtHgL33lasEh5JG7U6u5BIWuE4S1nImx7C4hQX/KR2JCl1Dy7A2/PhhnVDTnxaRS+miBiG4UWEfqDBfFC20CW16hIulUAk1tAx4PBkDXpxvhBQsEaMWoXkT8WMfzcFmKRPgxGYQu/MmYBCr4Ex/rEVWvk3tVNTjmffTYWM63eRn6mYD6c5EWtXE4HDUzQ0kns7KfVwP/UvQw6M2CXhZXzjfGreX8Wa0SgDP8b1efMNFLX3OYNoFSOYImIYhhem9GUoQGJmn9DWxHz1PJo3VeW6o5N63bXflZVywX0dX+jaoXB8ACkwvO7JEIGQm+GxyYP0iXX9j4rP2/MRUMINHQxkpXz4DcTB8nZuWdZBzYe7GUH6xSZBA9FvNBgB6XWoBgsyh+ZshV+/R8VhRCQzzdi5UY8zXoaiN55Gq16up9v/g53XMR3SYFs3K0vFCKaIGIbhBYVegwifhKC3ckHhMquB135fxfpExGu06grGBK2LW8cH5bp8SMPDQDl57iUKivlc2TRgjnVeFhI9oMEcFlO3s8kpMJMEyU50qDD/fj2vfrydm3+wgFrekf1qYrOYUo9BeYvxiyIzS1HYp1H/cUggj51NuNFeR6RLQzkwyTFuQUAVawbKi4jfliHw/sQ6OTNFxDCMGBIzllL3n1GEfE8jPPDH2qil5q2xOKZnFYecFVyHYXJKVoQOapmDLyMGmagk/tdN41RxVEDJ276dW5xL7KJbG7eMq3Fwdn0hiXNG9fxqWhfrCI6zs0j86Z0D74jt0yHTPe4e51PT8ZtAL8k9us+DGxpo/etViEVYIsMyqLwfGenPJCSOcXTGW62nqSU7EBsroIzFpIG7vxgjZ2CKiGEYXtDKyrSwBmWzL6NAvDsT6YFZfuJce6tYV7jxexSOV6qYA+6vbT61RgwvUTlGjFRxjJB93mkrBEqNcyDluqqLbvhtUedcVYKdEXLAYZKd3xgv8mvnc0SFCM6WAcrIiUivmBkUXmYGHFy6jlffN8HtSLzG3Q6kh9/AV3FO+FwPo4oj4JzZS1kfOfSwirmw1LxBgHuITQMQO9xniohhGDGEmW7iZbSsLxB6H4Je2wAUP5sJB4vqdgVSSf9UQ3lwQ0CSY0joFBNU/AtUHAZKS9JVsipOWKA0nt7OzRd7tt+1ccsjuFHnVXgH/Xk+z/TLgFs+viBIp8XIL2idl+cbBOd9S9aOpPmbGXHiEfN+ZPLPFhr0koiZUg6lLuLxGAnjHR7EOeZ10MFxz25cSBMVWpYgtnwwRcQwDC8hBcPlQ6YSlsxQBhpkPnj3dRz+XAX8iEtf92KM0a7tVVkpD5ac2TMBhc42JGLdL0I/0eBEpw03/RHP9l5UcLvobzw4d91ylDOY1IiCdFLPzSTBXsgsF+G/sTNwPOAy+AtQiiZcekMRiBg3DxESx/YiLaQ9XO9YxMke0u4BNy1WxWFCoeg6TRm4IaTBY1SoEUGsLQzSJzYNTBExDKMkktTTieLxOypGQIv22DZqVU+k8mx2XydIZkVYD7yXW9tGBTeO7akQ5KFZlLpZxUkH7v9pKBVzy/Xr4dyBd0nPJ0Qa34CDfAUVa8mL3EEb+WgbtZTkeK5ewP29QoMewli7mfW0rfPJE3X6JuRWvY6wlHqco7dhu5EcmGo6PMMUbq3BCEif2DQwRcQwjJKZLjs6D5ixC98xywWHU7Nz9V1Xish0CtxU5JdVzCGgIOOczTk5ywwP5CP8Y/U2WjVwLf8RCWeVsqkBaA1wq87KhVDidodS4fPqWRJddNsTUGK/Okt6XkcSHohj/gr3v06/jofl3Im0pD6Up1hbnnScXYeD2WOkSmGaOH4VZfE768M1vNPZm6g47gSUjk0DXNsjGoxgiohhGCXj3HEPiGt1xY7/b5lm+hkKnUJGa+NOxulU3NRSpowiEnC4MCOPIFthhr9VsWowSX8X9a4qZauZbYrwF5KSOh3XEGv7UA7ZRe56u3HMD0Ih2Rm16Rk4Sc5CeCNxtiohTS95peRa4px1CbEzLo4gJM/eSCnv+5I1lpaInxG8P0EDy2Pt3CK+jdnvxdSRYKmZy3dm3k+DUSS4T0MRTBExjJogsYuADVKi7PcShZ33P6jwqjLddCQoVJ8Khd+Pc8bcAzv/HO/KhuuIWJ8iMtst94+K5EiNGsml1aqI6wER/nhujws7I2ToW1GE6YsuXVSsKi5Nk5S6ICG0F/JRbG9LwMEcDdY1O9GMZigP3t4bxLshLW8aMzlPqlzVRSSRj2s5FfoQ3XsY7NZABFNEDKMGMAWxRn8NFG6pwZJB63EbDeYg1DAmi7MtpZ7lqKnihwyYzkYF43oh6oYuSt2Fa/L4wOCGAUqfgAojsiZIIOGmYqQ6xD+HelqyW/cipEnMLAzafpBLGxZyQyjO5f/IzTc1Oh/nCCwtfW3IqU9oVA6ovd+owboGSttpGowgQtdpMMoY2HTgue20nla7NWvGlezUa/YOCyGP3Z2kFf+nYgRTRAyjBgilnTtkLyFxWX4j5lGTM3SLTA9048xrx9BWo4t6zsM5FquYB78aylHdTedlYa/RKrNb/TVCj7qJrynOYyYq9q1HbtVs8aYl+BxyC7JKFCb5ZBu17qliLCHNOIg5+NfILeDwJv26IMvorpeQj+I8ssYaP9YL7dSyP67/cBXzeb6fpmVctOfTRi3NXpukajDOLt9dfhQO3LRrPxL8SkNeTBExjBqAgquA8RqV1ZppoGAudmgI5YJK5P5u6o4sklVF0NAJTyl4L3XGBppyua/SdfYIGhxGhH+swZqynqdcgYr9+ZHbBnqyahXYMup+HK32C1TMwy1lH1boxp/3dD0jKhQECmKMrUi8DUk9cAjttw0u/lLkH79iKHThzXSzV8nD/3wL3AFJYbullA3vnreXAfFHLKADCvouqSbraPXZKIBmq5gDruXJNUS/VtGLKSKGUQO6qPffeEW9S63jhT52Ph30WhUL4nwQCNOZKubjXYCqmrixfhY6GvdSjq+ImnEL3fIiCsZSltF/ajWt8bZkxxmnYMau0VEtptPAt+MqNVzCnA5qqsR9OENpim8lj4RlXw3lgJt/SoN1xzw66A2NPGUlrvINGpUDtPRHiNZ6Fbys8z8+SsURyKNJSbVim1vKFgidoX/MAek2I6RGz/Gry5ehQ7Rxy3fwoBH0g/LhjGJLI5giYhg1QuLXjZgWcHrxPGrdWSO8zKW5m63jV/4Whc7bNCoX4d9paExJUs8f0aL+rIp1D1rfXqPVHIR+VsiV+XixgJp2x/PdXsUx41q6Y71I6BueyiDMP6hwpeVjnKdUV2GpHMG5lcdTOUzFXITu0FDd0E6tr2znlnMbePBePJuIXVEWGcQ7cVKSVnmnKa8jeR/+6xlOZeesDDphabxEvAQ/9/cacThmwzNuKMY9t7u55U+4D68y5IAydg3Kh9hpyEOYImIYNSIt/D2UOHH+FPZOcPjXNm7+bjs1H9JBza93tgHtNHsXN7bcRs1nT+G+v6EQOFp/n4dc5xQEFcacLkq5YYyIO+p6BOlyO3Z+h1BZBgIajHGAVh3w3Ke3UevCYlvIwcf0L2POUuq9HNcVs3oqvyFB6z6tQpnwF1Zxy0rk32MW0AE7uZhWam1wK+9CSfkBmu9OYUZWjsIU5Ky1Ml7gvHM3Pofmk/AOfhbKx/fxPt4lHLqpuGfhV5tlfx1BRPi0gg7gYoZlAkmXtbqz62nAM4sxhuVWV26oUDZ4IChzhtKg6RikwyewfaWNW65dz08/gx8swW/8jaAsd86Q4IMaLkjk4SOx/4zd3i78svD0aq02OZGApncOUuarKk5KkLnv75KeokZqmyIobNx6GxnfEmjV7NZFPd4hlGqAc30V56qywyp5cVAaZi6j5f/SCC849+049/Z41pujIPC1ul23eMZVO8qCtxYrC9yYdJoTd+GYsfYLLPymTup+SMUcUCk5F/Fu+l8DfvnqTGSU/+L+nHGjdEmq4kIW7/in8I5fqGIOaMVdiWPnrEFTDvOpad+Ag0wrENfp1q6pulMulvSenbTifhUz4HmmkPY7FXiejuFnGkri4KW0/FEXHqKDWvcTlpUIRuoG3E0fojOzW3COp1A+NGWiwXxqbgs4roevMnCOq3GOHJfl7dTahJb+kMK7A64n4lYdzy+Ni388KwRnJ6k7xyU6KtL78P10/G5LpNcrNbpKyCAu4JNJSv1MI4ZBA+ILzDRUMbu8m5PGuN8QETlOv/ql/10+t/o41hE4VmbYp9B94Lun8Z1zTPdvN5Sj0UjH5jOI+aMqRq5ltOC8vSg8F7ihW40qiPWIGEYNmSWpr7oCV8VRg2P1iwRHFVNCHCignB3KriiB4iott1qmMzbcdTvqK1pQLaGVLwcSHoWr8BvnFcdVLLimWCXEsbX7Da59VCvSJoh/i7Ty9kYFIhdpsCKYEm7V4ky6IdHG0TMou7ViCj1Px/AzTVAYWVMHSqIbConpSmc3NJP5L7YC69JUhWfSwpHhPs6sCJ3JI27zru2CvJHY+BvvVHhX8SKdqquEID89xkKH+pQQB7O4KfZD6Rd5nxDh6uOh7zNbQA1xs6PcsE7mN4XuA985I+xdhdh5O94Is5uNNHSeou92qThlCp8XzpDn3l2qEuIwRcQwaojzNjlDnnWeSp3dAt7jykEr5BERmdtF3bdo1LjjWuksga43U78spu4XUKldq+JI7u2k3pjhicmBSMOZ+KxUmRyBXIc8Geu4rwDPh8KHuNk8Ktc1eGnX4/M7fTJlr05K3abRkw6kwwqUP81JSZ2+iB7o1+iSMEXEMGqMe2nx8n4EL7Ez2POvElsAFPYvoBC4ICHpt5W7MNlYgFb1lbiq4gahtUaC6DWK1MWU3Vri1oYh4fNVrJiE0OmByHuQqGUoFLIEihDycbczEahr8M7dI0JnsvAueH/PdDOy9KtJA8qe57C7hCVsckN1KH/csF7ZRLpkzEYkM35sNiJmIzIuNiI+DqPWWQHLkWix74tX1K1fkde9LM43yD/xjO5AhfEHprXXxlnnFwL3eTUKAK9H1nwGZbP2ZbSs5Cm6rdQ6bXOW3+AettWoDIGkT1pCK59UMYfseyf7q1gQ3LuzEXGu5EcF0uAKpEFm5VMcc4Bl3VGVpOVI2mj23sxBjF+O6hBKw4fy7TtwL5fjXkqeYVP4WcycITzjKhzPDTNFQFo9M9KOJmsXw99QEfCzSek51oXc7K4p1H8ic9iO/x2Ad2vYbXy2x4TvR+gPUFp+WcyBnDNwDVnOVbEoLHIRjplju4I6DrKUMwPIvW8vC/GLOB7SXO4LKXEHlKW49Za8tFPzqahXyppSu0GmHuVTcDqoZa6wfE7FouDaH0d5Pmwg20HNHxSmHPubEkAa0HrkicdxwH8HxHdNpR3+4taf0u8rBsfMxRSRoQLRFBFTRJAONVBE8nG+QjbQjlsL9U0NSP7rbDH0K8OYcGSnAa/fbjr19z9PU55HHTOWTveMCYANzRhGneOGbjrptqfdWg2mhBgTHdez5uw/rqfb/2NKiOEwRcQwDMMwjJphiohhGIZhGDXDFBHDMAzDMGqGKSKGYRiGYdQMU0QMwzAMw6gZpogYhmEYhlEzTBExDMMwDKNmmCJiGIZhGEbNMEXEMAzDMIyaYYqIYRiGYRg1wxQRwzAMwzBqhikihmEYhmHUDFNEDMMwDMOoGaaIGIZhGIZRM0wRMQzDMAyjZpgiYhiGYRhGzTBFxDAMwzCMmsG6H6adW/6M3d4u/LLw9G7q3uDCk4l2ajkHKfNVFSclQnR/l/TsqeKkoo2blzBxhwuL0G5d1PPPzBdVoIOaZoeU2EnFCAElejrptqdVLEortU7b3B22BAJKr++k3qSKBZlPTfsyJXZRsSyEJL2Ueq5TcdR8GZd+DzXtKRQ04b3cHWfYQYi3YpI1+PoFEf57SHTDDdTz1+w/iBbSHlPW0XaHq5hhCvV3X0+3/0fFDPOpuSVBwatVLIhQGCYpdbWKZePu426avRfjPoR5dxS+r8NRp+Je0CCUPhTH/4cM9yDS7/bN6FV3L6JFaf1r2aAMOwzHRdbwM536ktfSHetVLMoRtP/2/TSlVcUiyGq8MykVCtJBLXND4leoWBZM6eeT1HurikVpo9Y52G2dlQpTzrtSDfC83oXnhfxQDrIW27NCDU+vpfDpatTVbdQyHym7mYpVpVCamiLiwRQRZG9TRMZEEWnn5sV47RaoGMFVQqtlXesqWjWgUQWZT607BixPqVgEeTwpqdeoUBCUA5dhd1xWKg/knfXIOzNUrJhM5ccNp6GQOglp9lqNLsQ/8MC+AGXhd23UNJMZ9f4IQqGDoSAtVzFDG7d04vjtKhYE9xXivhIqlszh1PyaNNNpqGjej3OVlP7gGZzvmoQM/mgJrfy7xpUM8tlDSLM3qBgBx/4l7uVDKhalnZreDX3pFhWLcWNSeuZpuCB41+7Bu/YOFcsC78ofuyS1j4pFwbn+iHPNUrEI8gTelZ1VGHNwbb/CtZ2iYgXIID7uQx67hYWuS1LP7dn48sB7/y/sds1K1Sa+/LGhGcOoI1AY7b8jb/ZNFScl82jeVNcY6OfGR5Ae55SohDh2I+ZrUQlfQpQouYIaK6AMbY2C/UdpZhTufGYZSohjO/z+tDQ3PIBjXLqADojtRasEHPuD7dQ6iorPqC+4Advb8Vw/i4e7Egr2vR3Ueqz7Ivt9fWOKiGHUGa4wmU8t71VxUoHKe9cEr1uJRPgq0mG6RpcJn0gcfkeFmoDndxAzu6GiT2BrzERWANLAldHHh9xwH9LmqGxsleDwJwuopaLeCKO+Qb55m7BcASX2duTF3TW6bjFFxDDqD9Rh8st5dFBs1/qmSBvN3ht3vhKF6Ds1qmK4gH3EWNNBzScHLDfhKkqyPSmRrZmDa9qo5QsqVwGelma5Zg7N2UojjE2PfZEX72mnlverXJeYImIY44hIeI4Iz2WRgt3iqEhfkeD0Nc4YVaO8bEbPPO+O5zYc/HManYecq78puTDCb781dFwh6dToWJyBKu7pPe73gUhJNhcjyShdnLgJ972DRsUga3GuO7B3xrCXI3wD9iXayHgQPoskPBD7EzUmghBd5H4jIgdpVCzt1HqCMP8KT7BBo7zgul/A1ovgtdiuwtaD8+QY0nqAgkrfbKPms1SOBc/iJPcscKKvaJQXpPcbp3HfL1SMpZ8G/zQiP/xEo3MQoVOzz5+KXt8QSNOPbTwu/UWjCyCrh36Pe3O9TaUjdFr2v/RZjfEw/K64YY1xIyFygTtvKHSkRkVAuj+Cz1uGNsh3Y/8o0i3Un8TA04Tp0lKG4ljCEzP3T/I1jYqA7/408jp8Wwl5OQc0PnJpN2NVM1YFyEhmrAqqbaw6xGHUtFuCgxKMEOXnSUmdqkJB4gwKGQpIJ3VfqWJZQBFq2JzlURQUxW0UUJknqftSlUrGKVtbcPhHXGlsfnP5kYXOJVr7+yStWqfRw6By3gfV9BdwnTkzZUbiM1YdAoX0W5BQf1MxFwn/J0m931cpFpQb+6Pc6Eaw0FDMjSjoL1hNa7rzDZJ1dtD+woGrYI/GFimfFdTfdCTy5fUqx9JGLUdAeSk6gwnHOx3Hu1DFguCYn8Yxf6DiMKjE9u6k3hKUiShZo+vwMRyloALngKLzziXUg8qwMtywWcB0m4o5jOZdqQbvoXdvO8gDz6qYh5yLsuCLKgzj/jNAgwczhyfjDuZrdAQoEGl8tHVR6kaNiqWdmo909lYq5hAKH7aUupep6KWNW85D5v28iooZqxrGBIU/4rr6VRh3tqT0vKgSIqs1kAuHsb0KhYCig9ZXQSXk/LTMmJmknit8SogDhetdUJxdj8zHIeIv40s7zZzhxuQR9CohuKCXXY+Rm03SRd23+GZFoeUToiJfgd8cC6WuBRVH3DRu1zPyyw46uEjvUengeOe7qeUq1gBxM6PylBB/PguZKspnmyq/p1ufQ55ahEq+DQol8g39n36VAxpXCXxcXs18E4tQSdO3hzBFxDDqHGH6KVqMb1dxXBHmD2hwGCG+FIUdWq+5IP5gN1VVxZLooNY3YvfprBTFdRFDwfjcMlrWp1EFQYH8U+y+lZXGD+HNzkJB7/UDgXt4gSQ8qJNSizWqKEnq7k1Lw2xfOitbCw9+W8PVoBHP+qp2an2lyuMKs0SUbSiV38L9R4YdEPf+mTSzYgPgTRnnvyUhg/sjje7XqByQR7cNeeB8FceMPpqyItMDUyKmiBhG3cPTAg6vmUf7bKkR44Lz44FzR+w90Oq6AZ/OviEHRnkyQKXboThCls/gf1NUzAEF2Q1dkvqyiqUja7+Bz2eywtjjjD2Z5JMqRggk+HgX9a5SsWSW0fJ/BRIeixtyPiIiIH2Ob6OWN6lYBXhnYblqIS0s21fKaIDy04QK8i0qDiEBNVyFdI1UqMgv2+9AMw5R0chjCa18Eu/oQigj3t5DpOBxC+iAN6swJtxCt7yIJ/jJUOjzQxu09fP06wimiBhGXSEx9ij8pgRP+7kK48IANR6PXV7LU17cjJ5dgULlDxqRA5fRbT6X5m6GiuYEFfOQPpbwoyqUhRu+YeFP4xjfHrmFlPB2WY+WadTvhhViZp7IktHYHLihGhz7ZyrmoF3tH1GxXP6h+xxQyb97PT0dsUMYUziMGFFCybpbPQyPOp9NRpZSz99YyNtj5hoMaU64IcwxxfVO4jq+PbQlqcdr5OwwRcQw6gmhC9GS6VcpBxQgR7dRc3kzBUaBsPNomguurWsRPdCfoEHnqhliLrjGPQ6j1pK8V06lDe34R4w7aV6UpBUVKw6u8k9K6vMjN9fDoF9XFWGJn2Ehwai7wVnke3Hd3FDkjsnsykUk9rpwP+eoO/QxxymjQrxQxY0I/97tQgmWZuQ8cMMLDqfWilzDTxaYGi5GQnqHNJF+78vu6gNTRAyjjkCF81x2dkgMzN9168CoNGZk15qht6m4EQkyFYTr/oVwbyYuj4DDmF6OPJgO1lAEVL6/1mBdoz443pWVIvzb2XpouGI6KfUwnkWPinnwzvOpJX9YoyhCgZsCvUjFHLI9LXJ5Gx1cTT8oXqZQ39G4ty1UHAYabmZGUEjTlyP8ciYyB542SOIqUyOGbI+Sv+cS8TseRi11MyvSFBEPaGW4ltjKybyhcKh4epxROQEK2JeJvwmFxJv+eC5TmINFY21UiHN4fA5IX4IGhqftoSUbtyjYcW7ROQ0XgL0KlesRmkoDFa2VMd5Mp76ZmYrbi8RUAhUQMxTmSJCUrZiGbkRJAtc977WlQT7bnnjg2tKe4yhgnz8d+acbWnAhNVL2L2zHccN6xhAizueOH1T+cQr0uIP8lov5ETEmOzX1IyLysSSlfuaGNxIcojL2+1WAonLDLEm1uymfGkXV8iNyFO03fT03Pok0yO/67kpKz7Dxage17icsXoUB6XY40m2Jij64jVvWoADyLI4nf05Kalxdj1fqR6Sdmk+F1ua14cBDOq3QuHg5zKfmtoDjFD+/fwlHrB8R4WbXW9NGzccxs1vgMAb5Lo6d4wCsWn5E9B14EMG8eki+7YbSVHBp/CGksdfpWlrCN99AvV57lzg2NT8ihSjkDwSch/fZ66m3yP9KZlBmTCtlxpv1iBhGHSHEU93+Buq+G0KsO28oCYetouYquvveyDpqONKjhDjlIqdCeyd13yVET6qYSxFjwsOpdSu/EuLSYPxmvFSB7XQfAc+yivchMZVThgp6x9KZfNZFqcuR3r/JRHnh/8lUSmNAkO11izSG8/MZU2NnnI1MUOYsrclGSPy8Bj3IthqoOaaIGEZ9kakgHElKXYBKItb3hDB9JdMLUmXQ8vb4DqH1fTTldypm0N6Yq7NShA7XutNwhEFKxBipZmqmNRqsf7jQwnzhSxoYNUj//2rQQ5zBbzxCwXA+Y1l7GmLuUzEf6KSUs+4Rng8uZ3Q4j73Y+ZTVfzjndBrOoLNnvD0YuBhnUB1RZowsAcmLGvRQfr4ZK0wRMYw6AiX8yEJVGoRPxs472wM1REI4uGLIqBAtcI9RX3kcTq2vw3FaVRyBXJ/xDZCPhJdrKAfcxJQBGow1JkzQhtjKDOefOIuwCcd2OyMNqrjwXiLWh4xTEjVYMuLUWMVNdxZht9pzjOLEWyU4fZ0bsnMSzjdqRdHvsRfHFrpEg7kIxeQzfl0btTSpaESQiCHwEHiOdWN2YYqIYdQxi6n7BZQYqNDFW2igMN+eePC32WDB7vuSGOTwFBwoUi6IBN41ZJyjLhRoD6iYS4HZM/00GFuZMUnscEf9Ic9pIAIUqqp1fQcUxh4L6VWg+700nB0UC31YxQjIE2/bwI0ZvxQIj3rISZh9C7BJA7Fzkx+hn6Zdi6/XqpgDm9FqLL4h1hEU6GUbX5CncjFj1Ywx1nykTMZYcdIi8mQXpb6u0qSilsaqzgOhc/6j4jAFjSIdIh8dpL4rG3hapNeiVAM8t+ja3dziporuolEZoGi8jOPPQcjvy4Ld+i7R4RxHIWNCpPN//QWlDE6XYOtF1D1uQzSVGqvOp9Z5AWdWAPYgP0pK6lMqjAo8/zOQ0N9RMQcR/kgXdXuNOeOMVUOR9qWU6lJxmHZu/iGepddLLPIB/iYHBcQDKB/d7LocSjVWnUdN2yHvP458ljcjR/6K4/sUlCwcfBf/8fXWvThNBl51Ld1RUs/QJDNW/Tzyjd+jqRrGq5RDIWPVUha9KxfrEfGAzD4L20cn84ZM6Fb/NOoEFBgXC0n8yraopKbS1K1QWZTdTT/EKmqeg2efo4Q4ELcFlI07maGneDb8wquEOAIOCrVW/RU/ccMGCvdRoa6ZSv33YIdk93Kg7kcP+4bLsjCFZbuPj2O6POdmyNyZlXJBPggC5l8wpUdl+9JAwQk4lmdaMO/ly1/Dm1cJcfBW66khdtXlyYww76fBCMi0Lu/WBaaIGMYEYboMnIqd14kY2HKQ+ftM8pDK5ePvLh8VuJ4TXE+LinnwCg1EQAE6IVx4X0+3/wdX611gDPe3ZwfNfqsKFeN6EFBpeJ2/QTl9bjrtGJcnyibjNVdkIY4bN8y3W8jBsbger/ffUvB57B0tUJRteCaPBXSAsw+JcRooL66hwBQRwzDKI9P1LHwMKgnvkAVajEcKcUVLfDt32ahcxqBVya+9m1pbVMhBRDo1GAHX8r5DqflVKlaELuyFZBljJHbmEKPS/oyGKyZBfBpuwjs7B/G/W0SLSl7ltBQWU+oxFj4RzyCy8q0D53QLFVY0ZN9Gzfvgv1GPvaPnkPnUuqOGx5w2br6wnZsvHrF9U7+qG9KUOAlp7TVWRTlxfTd1exdTrAWmiBjGBCJJ3Q8Gwh9UMQIKnu01WBaDFB6P/xaYilo5HGO0+i5K9aKye1jFHNy1NDL9WMWycTY4aU7c3cYtN7dR064aPSYINf66QA/BiW3U0qzhsnHKFDPlOBQbCZS5MVkIMUk9zu4lZqkBnoaPylaCHoNetyzcEJCMi08Rp/Aw8adwzo8MbajY49cbqgEZpYyd2ZcflCEXa3DMcD0ybs2ikZt7L/XrHEwRMYwJRif1XIOK7yIVqwKzmyYcYQBV3XewDa9gW2Tz+qJAIb2wnWZGnJc5PyQiodcINAu/t52a/p8KJTOP9tkywXw1KovNodC8G/f2F+cNNH6IaHR00W1PYOddGwfX4FbIvWwete6sUSUDBWrrkBNX4Shx/h663KwlDVedWdLzZSG5WcVRk53+m1mkL5+nRuShUrbV+r9cxmlF3gTJXhqsS9y7xiyXI+/FObq7sZO6Y12/Vwso6O24jptHbgGzG16OYIqIh+n07LcGZcNWk3nrl6ljvrCaUTmrZa2roCMzFyrhMGpxBevMrLQRVELLkpI607nbLmUTYe+Krq57WGjGESrmENLmaNFL/KwkDr7bTs1fWkgLY9ZzycW1BBM89Ub88e0aBXgz55J8FbekRjvcEwcLf8nZa6iYA+5/lwTLyg5qKnlIYj4d9FooUD259zES6ROhspW0cnCK4hQZPB7ncorWqInz2At+48tPcRt+76ar+9i7jWZnZnyOJWgE1K0i0k6zd0F2X448F2cb0hfI4KdVGFOEw5J7AnG9udj0XWOyU4/Td30cTs2vGWS6p0DLJ0OxKYluvBvHiE4zFTkmSak4+4cIbkn3qbzBrfjpa8HfmJSeeRrOoZ2aWoX5VlxDrLKBwv8efJy9mtbevIpWDWj0MAupdfP1JCeiRPsKxDg/JHc9JWubfP+vdPruSHAf74Hi5KbKRspVZQDKyk9Zwu8maYVbWDPCEbT/9v3U8HFiPgMH8brAd4QSfmop9f5IxVjKnb7ro51a9scddSPYmI3xU2z6bhu34BlHK0gosHt1UXecZ9cIzgBYOBHz++jaOPmMdvou7uPXuI+cHkQ3xNglPaMeAqxw+i67lXRxT6cwyccguqEzLyjPTkV5VnQ4rxrTd5FO9yGdcoy1kf+/1yWpiN1U5IUxRcSY7IylIuIqzHWUbgwp8Qa0kv+o0cOICBqicqELr6HEy8UMytxiaGg543rjezd9hatTGqbQ+ilpooYE899wv3kOs2RtIOnd0hQMTwfuol6vA6RWap22OaUz9iW4ll/iM9L7gQIonZD0W3G8/yQoHFxCK3O8wBbyk5HH8yj0XU/QEyh0X0B4C5zvLQjvi32sy2r87j8Nwvsupu5HNCqDG8YhmjK1gYI3ojT09zCJnD1IcnEDJSRJ3QWdxqHSPgfHwTMsTEaxymzyJJ4PynZyhpZoact+hRQyB/77S1R6H1LRyxyas9VU2hDgt20Be5zRQcnEM7nJBfFcX3AxLhxHKc/Hp4hkhsloCu4n8Wr84N78fIqTPgBFb9gzapr608voLu/04KE8mxGYVyKd3pIJ5yCrG2TKngPUFzKt73NeY/WL4WthCprwjGIWZBQXHzOtPIuzB8F95PvaGZUisvF9bNymgdPemW94XrfjvD3ZMCMd5JWQndL9diRICYa68h0oMmeq4GUo3+AxdUCB9Xq5xXVcg/N6bbuGwPVBL8rYNmG3EVNEDKNExlIRaefmxXjtFqhYELRaW9FqzRQ8hcha7HP8AnkeRQTvuZN94/VxDCSlx7skPCqpeKdJHlBo34pCe46KwxS7j0pB4bcGaXBIknoiKwWj1daJQnB4ReFC4LpDXHfRIaKxug/lkumywweLzZTBNaAy4+H1YQrRII2v/D3dGushVmEcEy3kjCt4Lz5FBP+BzCUPZeBZ/QkV1TtVzAHv5W/xXpYxTTfXoRz+/0f8f5aKVWW0igiu7Ve4tjEy5HXJSl9B/v+ayrGgXHDLSYyZcXecImI2IoYxwZkpqbPxglfNqLBWZLqdxXUt+93ZVwLS5TkRmeNTQsYKdx9QYD+Iymm4NT56ZBBpcxaUwVOqPV23RCQhgZut9e+saEwEkAefRF5sL0UJqSXWI+LBjd8HFHi6/SYT/FIXdf9BhUnFROsRcXTQwTuEPHgPXujIQmITpUdkCGfUGTJfhGcwysXMZPmgBCcuo+7HNSLCWPSIDKF2QD9E8NBsTMXcmRb+xA3UfbfKRRmDHpEMGWNQTrghgshUb+sRqZ8eEaSjG277YULSF+QPgxbCekTqiAQFxzDT4sm8oVSp2IeDMf64pdJZ6CgEI4aYEw1XmaGwakGpNR8F/E0ovMrpAYDuKLfj43BUQgcXUkLGGrfGTsZAV8IDcR+LsZXsjTR7z3ILKvcOHGP/cpSQsaSLVtzLImM6W8eoDOSvl5FvOkn4A/0ybWe8Q18uRwmpJdYj4qFUg7NNGWTq+6Hh76nipGIse0TaqPkkKHq7q1gQKBYXd1KqoFFYPm3UBCU6d8qnSHilq0BUzNBBrccKh5n3vETSmaETD23UeiBz6J0R40XooSSl/leloriZJH2UOCRgPkCI38wkr0G0M1LF43FL17vVZ/l+3Oi9aFl1lZNmeNffj3e9pN5PQUJ2UcrNyqmIQ2i/baZQ41ycbzbuYw/cxy64/i2R13DZ4gw0H0P835E+K7G/aSl1+/1llEDWuJRKW/1X1n1tpFFnKTi/LMySMxU6lOAH+deM/H468nvpHk+Fn0xSj+tFioB7Ohr39A4ViyLCK/HuDhul4lo+gWsp259LKeBcz+NcpRhbe8F7exTe25J7a5A/kB3pBVTgG5jCR5GX/jmNdnpwtMN2eK6fw3PdRsWqEwqlfLO1TBHxYIqIy+imiLhwtRURwzAMIxcbmjEMwzAMo2ZYj4gHdUx0pIqTlceS0jNWUxDrGusRMQzDGD9METGMPEwRMQzDGD9saMYwDMMwjJphiohhGIZhGDXDFBHDMAzDMGqGKSKGYRjGpGMezZvazs1r2rlFPFu/84yrPzXGGDNW9aBOjk5UcVIiRI90Sc9HVZxUmLGqYWz6tFHrHOa4NZrkwqSkTlfBGGNMEfFgDs0yiog5NAOmiBjGpkkbt5yPCtAtVZ+DkLzAErwpSd3PapQxxtjQjGEYhjHpYBLvsgQs9E1TQsYXU0QMwzCMSUUbHfxqqBxvVXEYIXp4kDb7kYrGOGFDMx7cIl5Qi9+t4uRE5OkuSk3KFXhtaMYwNm3aqOUDzPRLFTcickySUlerZIwTpogYRh6miBjGpg3e8avxjr9PxSHuTErP/thLVjTGCxuaMQzDMCYNC2lhArtoj7ewM1w1JaQGmCJiGIZhTBr66Ik3MvEjCK4a2oTkJ0nq7nXfG+OPDc0YRh42NGMYhjF+WI+IYRiGYRg1w3pEPLRR68eYw0+qOFn5Z1JSh2t4UmE9IoZhGOOHKSIezLNqxmLLPKuCaisibdR8EjPtrmL1kXVfS9KqdSplgGK9JxTr41UsCSEewMezCD0bED0lFPxlNE6eyr4GoWuSlLpHpQjt1Hwq3tHXq1gQluD3ndR9h4olM5fmbtZIfYcySyuOgndBdkG+2Ar7NNLnRSZ5OBT+AwrR60bmETzj0/GMd1TR5aHlXZS6UcWidFDrfsLhe1QsG5HGH3XRbU+oOCoW0AFbpClAGgRIA3F+N3ZB2bAV7nkQ+xew/zee1XIhQRr0/jv7r8zzOQPPZ1sVnZOwmzopdZuKGTqo+eCQuV3FoogkLlxKyx9VsSyy99E4D3mwGeJe2N6Iq9oM192IZzqAe3kSz/MfuJe7EH9rknr+mPljBXRQy/uE5Z0qRhFejOPfrlJRWqm1YQsOv6FiEbg/KT3nqDBhMEXEgykieFdMERkTRaSdmxfjtVugYtXpl/5tb6I7nlcxAyqFI4n5WhUrBnniARQYiwMZvGQJrfy7RpdEudcgIiej8r5ExQgop7qxa8lKRRA6DQX/T1QqChSBN0IROAP3+37kg801ugiyhCT85MvU8OTmLC8hnabrFzi9fLNLUl9SsSjZHlkp+XojiMwspMSVQju1voUyacDH4F5maHRBkF6/S8jgp16kxue3QBogqjH7DRA5G9eUU5kiT3weeeI8FYvCwvuXq1AuoJZ3hEyfwbW9d+QzKQEoWPLTQer7+TK6y91LySBv/ga7k7JSFOSHpxOSfifeoSc1qiBQRKYhPderWBAcew3y2hYqThjMRsQDMuyDyLiLJ/OGLH2rJodhZMB7sQd2Z4Xc8DcUtsk2mp1psGwquAIfiuK30Jp9AHf7kdKVEAcvEA7u2ZLC48us8OqKdpo5o41bvgcl5K+4pw+UqoQ48Nsj09zwJ6TBcRA3KiE1oIMO3gF59DIoIW5WzHEVPJNdoSSd38DTHuqg5g9CxiGqA/LVDkinK1xPh0ZNeiKJaz0ixmTHekRKA62vNAqQC6bLc+csogf6NdpLvfeIzKeDXss8eD2e+zs0qlKgx+eWqxOlR0R7gq7H5Y+2JzSSBuPZI9JBLXOhgFyGC9heo0YNnmFngwQnLqbuFzQqlmI9IkMgkb7VJT1nqRiL9YgYhmHEgEo7gc8z1/Mrb34PvXvYHmCisYAOeHPAg71VUEIckcbdRMDZ8EAJSeHyqzEcW7M0aKfWE4Spq5pKiMM1TNJIn0Op+VUaNWpwjWe2U1PFtkCbEqaIGMY4Ml2C40TCbbDN0igPshzbtwtv1JP9bXGm03Odes5t0GK6W6PzeSbvHBci7jJsKxEezPwinpYB7r/5cGp9hcoRRl4DxLuysbmghfjA0G9m0HNXarSXQAbbM8eScI5G5SGPDh1rkGZE1xRR5lPrjmlOLEO1sLNGeUG6PYLPXyDwFWynidAXEXcprrmkcf5ySNP0Xw1du4j8P40uSCjhCUP/mU473qvRJTGPWndmDl0aDBvY+pF/Yfs57v/LLg2wuV4e5BFZnf2+PF6m4IJp0j/DbRBzevGGQPr+fug3ndR9p0Z7cb1uUKZ+jWCRYSH5K57dT5C25+A5ng75XMjX4Iv/Zr+Pg/dsZLqpjZq21ggvfTIV+cM9C9pNo+KAfhP8Cu/N61T20k3dfUPPNhQ6WKNzwLk+675vkOA1GjWhiGiuNjRjTHbGY/puG7W8iZn+oWIuImckKXWBSl7w/wX4/2IVh/ENzYwE93Y77m0/FUcgf01K6m0q5DCP9tkyoCkLAg7OhliocL1xpvTM/yrqRZW9oIxZgd0BWWkk8dcQRxs178PMkQoKFdjDXdKzq4pevgx95m5uuQmFYOwClzjOPSzh/ySp1yl+EHPJHINaj0RF/r04ZQaVXFlDMyNp45ZeXN9sFQtxbVJ6Fmq4ZJy783W8+g/IE00a5eNOpMFnOqnXPbcI7hjr6emjUZu4POvvMfAMzYwEeeI57JySmodchzxxpAqxLKCmPUJmN+NlM43KB6+yXEnCX4l7n2fSzMYdafrhbrgI6fFGjfYgS3BNricjkh9GcgTtv/0ANz6tYiHuhaK1/7V0R9Hhl8OodVaCJTKjR4Q/3kXdP1VxwmE9IoYxARmg/l6UggUr/GrhZg0spd7L0DB5K85ZaBrhoauoBS3MicHd1PKhwkqIfG+N8L5QQpw9irfScUoXKoBFjTI4Ez94QKOrghsywvXlKGw4h6ushqfJbkQ6irXUfayjpz9RWAmRc6fLDrPjlBDHIlqUTlLPFYEMzsLvq660F8MpQiEHv41TQvAc14jIYVAGjyvUqFhFqwa6qPfaGfKcm6Yca5+E8yxAQ+AUFarB3hu48bsanpRYj4gHvNBHMQdHqzgpwcv7f3hxP6PipGIi9Ig42rnZzWzIGdMfix6RfNqp5VMoOdzQTYRsoR+8aSl1x3bX10OPiPMRMpX7HkZwu2xMhB8npacsp4aHUdNuAQd/RaE6RaMyIE0q6hFxM3jwfM9UMQOOdQPinsA5PqRRGxH5GPLNz1QqihtKS7M4pcarwCANz0cafk7FkjiMWvZCi/0eXGPujJAx7BFpo9YPM8vPVcxDNoQiB0KRLjis44HbuOWnSOdTVc7nmenCuy6i7jUqRyijRyQDypqTUNZAoYrHekQmEUyBm6Z41OTe+FDsjTqGRY5DATR35PYcDbysX48ZaP3+EJXU71TMAUrO5qgU/kfFumUqbTgZuxglRP76lKwt+x5uoN5/MEmBlnTpuKmdQnyCisOgZX8DPvyLs3H094VIU+iUmRglRO6eITsUndGRzw3U45Tjq1Qcc1xvCHOYo6yNBJX7GRUoIQ5Jy4xPY+ca5j62W0cSp6RUhJsp1UGzXW/MpMMUEcOYoHRS71/QCrpl5Oa6l/XrMaUhY0TpN2JFK/JkN96uYn3Cnh6FIUQ+WXE6SvgNKITvG7mRBGVXzFtSeh7ScScVh0FcVwOR81AKXTDC/q6nTcNFEeYCaUCnuSEXlcoCCvI5+WkQEnsV19GyjlbPwRnfoGIecl8XpS5SoWyW0bI+5IX4oUaWj2qoLKDkxazyy5sJJ65zXmA1YtJgioiHBmq8CG/p7pN5C0QyQxOG4WMxpR7DbklWirDdTrTZgRquO+bRQai4+O0q5nOv2oRURJJW/B8UwkV52336dclASfiABkcgf3Zu1DXt/5KNy4GZqaRekYzLfaI3q5gDKso7UIF7ZzaVQielHs5Pg6XU8zf9urowxxrospCzu/ApbCXj8gLSw+vu3Rm0VuTUT/hSfDyuUj67hdwwYYdYKgV5MRezETEmOxPFRqQSqmEjMkR23Rx2zps8yLk43hdVyKHWNiJI+w8g7b1TevG8z8Tz/o6KNcH5ZBnk/idQPE/VqCxCX0lST2bpCTzHb+A5etJXHp0pqV2Lzlyi5tNQif9YxRxCCT+1lHp/pOK4UKmNCP73L+wizxnKQ7pPpm57C93yokZVDNLqDKSVP09kZ1R9X6Uc4mxERPjj7BQRjlXk3Y+89j5mI2IYxvjA/Oo2appZyuZccuu/xp0EpQt5uYzrcag5zBLrwyW+23z8GKSBkyJKCGByXk+zsHCXBvPg166iJrewW0GEOTYNGiioeRqUwhyasxV2cQbJ91RDCXFAqY3tIUM6VuAEL5wKhbITee0KjYjC9AOndKi0yWOKiGHUH6czB3eXsglNH7uVfIswlV79EHberm9EFvThUUsk1qbAdSP0lT2MUm2Eo+7BUWk94myCVKSZ1HMn4rwzMoSDosMzaJHHpYE8QWtrngalMI36Yp8jE1VtKKifpt2vQR8F/I34wbVNy+wlcMawMbPLeGrAcg0aG2VPyZ6ImCJiGEZFZI0ZM6us+qjbAhSVsNf9N5Sn/nJXWq02maEmosjwFOKc589hskMv3KliDvjtwmI9ZVDGYlygy0vjZfA8eiQ2j+FZPqvBUXMz3bwWx4tzNlZ2PsexMopIkrqfFQmORUyc0ffr0di4NBvctIncoNmIELn1BAIKXq3ipKSBBtd30opCLYFNlprbiJSBSDiri3rdCqMlUU0bEQfKC+fi3ONNU17E8bwu32tuI8ItD6LgixhqSh0sGKa+KyKzMdJCb8tOjd1IG7UeyOyWA/AgdJxzMqZSBJznUZxnFxVH8kxSeqq6TkspVGIjgndoPt6hmCEq+jru4xwNjxpc33+wi0z3dj1VyDOvVzGHeD8iufZTuI8v4D6+qWKEUOjzS6nHLeuA3zbNdD2hmS9GYDYimyCNxB92BkGTeQs5cbUmh2HEggrfq2ygxb1Wg/VITAuUN19Ie+Q4IxtPjqL9piPljlFxJKvylRDHLOruQfq7GTRRWAoOz0AJiVk/KNPLEGmg1ie8TgM+qtkjh/SI630ZfT5HQ+c8PMfIcg1DQNn8JpTOzJpKaWosuvrvRMQUEQ/IFC5zPTOZN7x5roVi1ITM1D7Xy1F0EwoKFcZjilueHPkElaeXqnWNjwGxhfla2qFqq6uWyzpqOBLKUESxEyHXPR9BZ8Z4fZRAETykjQ6O7dVFGReTBtyAlnyct9m6IkFp10vhBfdXtXs4hPaDEpLnKXYjJXtOLYCQhKfgmp2n3wjIEwkolpc4T7iDlKjn96piIpqvDc0Ykx2bvlsaccMiDhSqv+uSHuelN0Lth2aaL0cavF/FHFjo6E7qybHHGC/auOVWFMjR1VVFzoZi4a90Wd6Me/F7gS2Qj9q5+Tr8+QgVcxFa4GZ1qDQuVDI04xThzTlck6moI8i/8L+yDUl9tFPzIcR8o4o5IJ9djHzmdWxW6tDMEO3U8i48zxSeS2TGlGPoXMgnfcgnucsI2NCMYRiTE95fA1HEv9R/PYBC2+cMLIOw1MQRmy4F7z8389ehtF7s3eKUECDMzo29FyjYdZcG5eIaybj/GBfs/IbDqbkqS+IjHQ/SYIRApNAU9rKA8vdHPJfPqhgBysdHMnZBJE9p1CaDKSKGYVQEKkJvr4IDheUtGoyAll1ML6u/JViYuP9I7PALCj23pL8XXNuxtfDNMsjhKahoqloe43hvXUAtXj8XhdIAlfsJ82heBc+iFsTnszTxqFfIdUsVIC8fr2IOyCvhAAWx56+ELkr9WEi8Q3GAmcOfCPEjKm8ymCJiGEbZtFOL6w3ZJyvlggL670lK3aNiBFSQcWP7lUyF9NoC4ByxY/dZHxzkZvtEQCX8CqEZH1GxIjqoaXY7tb5FxaJ8GeUwrjfiO6QapGNcvj9J61Ko8OLsDbZroLWxvSml0E5NrR3UWpWhkUKwxDsFE6bT3CrLKlbETjTjOJxlZxVzwDPrXUbdca7aKyYh6dPwbB5UMQ92foPelQ1vOpgiYhgTFDfDwy0VP3LroOYP6tdjhmsto5CPXW6ehX6owRjEGUT72G4etXoL/TgClpkazAGtxic0GMEZeaISiV8ll+nrOlRSNs4uQJivJQ7/1EbNZ7lVdPWrWFZRs1u47bUqVpvjfAsQOl8hULouUzECnu+3FtABkUX3SsEZybo0CFnubaOWz7gVcvWrqqNO3ry9O3jG20/hvgtVLBt3H8go56sYRaRIPq+MJbTyZSgj74Gy7F1JG/dVM2/KYwXuKRczViWaT037QkPzGNNNJhLPJqk7rotwk2aiGKvOp+YDAmZn+DkMCq+buqTnUBUjjNZY1SkhCV53CQqOozUqB5z/4TXCexQqN3DvJ+Le/YpAgbU78nEV3Dpe/SDuJ9LyDiU8YSn1xla0zldQA/ND8YW63DcocvAy6o1TmiK46bcbuHEpitVhGwukxz0k6Q900Yp7NSoCnskVuIdjVRwGrf0PhRQ8qmJBmMO5+PycijmwhB2osJMqDtNOs3chDpAHY40j7xmQ/rk30R3Pa1RRFlLr5us4vAn3M9J+6E6W9AcL+SWqxFh1iIL+VEAla+dkPJoyd+XdxwjkzzMlNbPQmj7lGqvm00GtxwrH9/iMZKIbq5oi4iHrYEaq5gxnIoIW5f2o0LytzU2diaKItFPz51FYnqdihrFURNw1o8T4BQqNVo3KR5Be7UgvVMbxuAK6nxufwnEiPbK4/v80iMzSFWYLkn1Po46gBDU/S7Cj81ypUV7aqPlsZv6aih7kb7ihE0pxGKeKzSLc02yNGskA0uVopMvwWjFDuPVSpnK/S4ucadC4h4e6JLVbJlgC86hpuwYOXC9QpPcDx7oGx/Iqjijvz8XurKwUBSf/C5LzxEKK1BDOODTteoO8Q3bSR8JHJKnnBo3IYTSKiAP/d0rncVkpCu7jfCjI55RSny2gpj1C5qtRPe6pUTngWCHyRfNSSq3UKC+jVUQc7dx8Ma6j6FChKSKGsYkxYRQRbu7CKzxfxQwoJKuqiCygA7YIqeFAYXpv1mgv1p8CkO/gGGeqUJA2brkRhc8hKuaAe3gYrfgPJ6n3Vo3K4RDab5tGnvIl/P90iJEyDP+/GmngcwyWgxvaWsfb9iI9YsfcnVKDU/wGlehlzoFYfgvYDeFkDU3ZXcuW2dgI/x6UDe/wuY+PXQV3xEq7pYI07URitKs4AtmQkOBVi6k7YsCbnQIrd+J/BZRQGUTD5Je4pstnUc+K/DSYRwe9IcGDzjD000iHzbOxuTibB5Z1M5O0yuv3ZrSKiFPopnEfFMb4dYRwrEdxIecxNf6+k27LVxAYium78E4625gPF8znJT6baigirgeygde5Xs+CjUJTRAxjE2MsFZE2mo13q2E7lOWvRms8bgl9d76iXfKoHA7AC5wztOBTRDI9EDQlU9Ewhz/EZ2ShPFQUa3CszFREHDeA0oFr5FdCcvuidg74/6WzJHVyoa7qkbjZHGmmu3HOWDs13Mvf8bkSv3kKikCfsLjVVvfC9TQhLsaRmvSx0O6dlPI6h8pHhyfQsuWiSzrgetyY/WNIm6cRds7cdin2v8x/JA2NcWOPgus5GKQg42KeWZzCmSlvcxD5qFDglrjPkKb1d/kUmayTK8qu0srhYci3cdN5z0NldVtAYYi0uU3jMqgisQL/3UGjCiAvIn84A01UsDId/3G2NAWdwCFvvIBzH7CUenIWomujpl3x+LNGvUxXIT0j7vWRfivw4XptqIF4pU+ZGsIZCAuHKVwT8m08OCbyqPzD5SvcC9I0s/bQm4r9z4F7uQL5/IRC+byDZr8VyjvSJP0K5mCRRg+DY1xJEvzKhdOUvrfY8J8+H7wrUWd3Q5giYhibGGOpiLRz82K8dgtUrDooZCOKCFrdR6LGc13mYwLOef4s6fl8qUrIEGjBn48CKNZvQkWgAk9S6mKVSiJbgYnrofGtvVIxSJf1uJ4juiiV4wxrPjV9MuCgLENHEdkXx4n4ZplPzS0Bc+wy9VGkD63xzKJrIzmMWvaC9rkMaVCRgWocTsFFpduBSvIPGjWMb2ixECy8fyd1F/Tb4e4jweLsdMoyei6RS56StR8utigg6lDXwChpFhTKl8NRvixRMZYOamoXDtzvInW2Y6IrIjZrxjCMikBF+xgK0jYoPp8rVwlxuP9hFz97pTxQ59GXylVCHEnqfjAhg/vjELEGjxXwDEt4WL4SUq+4tWwaRPbDM80xfh4dshrKwyE+JWSscPcRSHpf3Id3WK8ypA/5/PSk9JxSq5WJnbEx7um7Km5ymCJiGEZZOAUETbAzpkv/m9GaK2iYWgRxhTsK+VMzLeeKkdW4qMOT1BO7gmkxltDKJ2dKag7uy7nrLnmmTAxdgQy+PUm9ZfRU1B5nILxG+EA3ywTPY5RrTcn1LI1Ig57bNWLccM9ylvQcQsIfQF71+ospAzzL9N7I524aMA5XO/BszsIlpFTcpLChGQ/zqfXtCQon5YyRIULi5/HyRaz8JwM2NJMLKqU0CooHcOwUhCVrKLgV5ULM6q2VkbXVSEAJkFOQPjtqdEFwPQ+w0EUBDV7qfC9o9KjJeladcaIwnYR8sC+ivN3hI8G1rMePulD5XZik7l6N9lKPQzP5ZKbhUngy7vxEpEGJDrRkLZKqk5EGxYZQHGMxNJNPZro5rTuO8SzxjJxtUSmN7+fxWyhS8pNkAcd8cYzF0MwQqJt2DFjcNeXY5ZiNyCaIGzMOSfZQcVISUPhi3KyFTZ2xVESc182QElUdhx+JUPj0UkrlOHhyTsISRPHrwngJX2KSDcgHq/tpy0eW0bI+/WJMcQ7ANiPaE+feBzXPW1FAvQKVwtbYO8Xnv9ieYQnuRW1y+2LqHnNX1x108A5pGpyN8+2F63GL6G2O63FGmi9i/wIqq38JBX+G5nLHIuouqVfHeRwNibyu1+MYoL5bff483LTdBCVKXhuGKQxRuf5OxZLITE0mmY1/74kD7Ipn4WbGTIX8UjYNCO9HeC90sdvjZsX4yKZDOeVsmOqiXpcHKgLl+itx/3j/6G1QTJzvGXcfztgbSiw/DWXv7wjfBUX7ntEo2m7xOska8RYlQQMrXA+OiiWhxrA56ZbGNS+j5cPGzRMNU0QMI4+xVEQMwzCMXMxGxDAMwzCMmmGKiGEYhmEYNcMUEcMwDMMwaoYpIoZhGIZh1AxTRAzDMAzDqBmmiBiGYRiGUTNMETEMwzAMo2aYImIYhmEYRs0wRcQwDMMwjJphiohhGIZhGDXDFBHDMAzDMGqGKSKGYRiGYdQMU0QMwzAMw6gZpogYhmEYhlEzTBExDMMwDKNmmCJiGIZhGEbNMEXEMAzDMIyaYYqIYRiGYRg1wxQRwzAMwzBqhikihmEYhmHUDFNEDMMwDMOoGaaIGIZhGIZRM0wRMQzDMAyjZpgiYhiGYRhGzTBFxDAMwzCMmmGKiGEYhmEYNcMUEcMwDMMwaoYpIoZhGIZh1AxTRAzDMAzDqBms+2HaueXP2O3twkJyJROHLmwYkwdpxauxcyYktFsX9fwzE20YhmFUnYKKiGFMdkwRMQzDGFtsaMYwDMMwjJoR6RE5jJp2C0imq2gYk5oZ9N+/LaIH+lU0DMMwqgrR/we2XZMZlM3ioAAAAABJRU5ErkJggg=="

st.set_page_config(
    page_title="Asistente de Escritura Académica - UNAE",
    page_icon="🎓",
    layout="centered",
)

# ---------- Estilos ----------
st.markdown(
    """
    <style>
    :root {
        --iavq-azul: #4c013e;
        --iavq-dorado: #9c3d7d;
    }
    .stApp {
        background: #f4f6fb;
    }
    .iavq-header {
        background: linear-gradient(135deg, var(--iavq-azul), #7a1a5f);
        padding: 28px 32px;
        border-radius: 14px;
        margin-bottom: 24px;
        color: white;
    }
    .iavq-header h1 {
        margin: 0;
        font-size: 1.7rem;
        color: white;
    }
    .iavq-header p {
        margin: 4px 0 0 0;
        opacity: 0.85;
        font-size: 0.95rem;
    }
    .qa-card {
        background: white;
        border: 1px solid #e3e7f0;
        border-left: 5px solid var(--iavq-dorado);
        border-radius: 10px;
        padding: 18px 20px;
        margin-bottom: 16px;
        box-shadow: 0 1px 3px rgba(0,0,0,0.04);
    }
    .qa-question {
        font-weight: 600;
        color: var(--iavq-azul);
        margin-bottom: 6px;
    }
    .qa-answer {
        color: #222;
        line-height: 1.5;
        white-space: pre-wrap;
    }
    div.stButton > button, div.stDownloadButton > button {
        background-color: var(--iavq-azul);
        color: white;
        border-radius: 8px;
        border: none;
    }
    div.stButton > button:hover, div.stDownloadButton > button:hover {
        background-color: var(--iavq-dorado);
        color: var(--iavq-azul);
    }
    </style>
    """,
    unsafe_allow_html=True,
)

st.markdown(
    f"""
    <div class="iavq-header" style="display:flex; align-items:center; gap:20px;">
        <img src="data:image/png;base64,{LOGO_UNAE_B64}" style="height:64px; background:white; padding:6px; border-radius:8px;" />
        <div>
            <h1>Asistente de Escritura Académica</h1>
            <p>Universidad Nacional de Educación — UNAE</p>
            <p>Proyecto de investigación · Piloto RAG con CEDIA (HPC AI Gateway)</p>
            <p style="font-size:0.75rem; opacity:0.7;">Versión MMCD 1.2 · 24/09/2026</p>
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)

# La API key se lee de los "Secrets" de Streamlit Cloud, nunca queda escrita en el código.
try:
    API_KEY = st.secrets["CEDIA_API_KEY"]
except Exception:
    API_KEY = None

if not API_KEY:
    st.error(
        "No se encontró la API key. Configúrela en Settings → Secrets de Streamlit "
        "Cloud como: CEDIA_API_KEY = \"su_clave_aqui\""
    )
    st.stop()

if "historial" not in st.session_state:
    st.session_state.historial = []
if "historial_retro" not in st.session_state:
    st.session_state.historial_retro = []


def _sanear(texto: str) -> str:
    """Las fuentes básicas del PDF solo soportan latin-1; se reemplazan
    comillas tipográficas, guiones largos, viñetas, etc. por equivalentes simples."""
    reemplazos = {
        "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
        "\u2013": "-", "\u2014": "-", "\u2022": "-", "\u2026": "...",
    }
    for viejo, nuevo in reemplazos.items():
        texto = texto.replace(viejo, nuevo)
    return texto.encode("latin-1", errors="replace").decode("latin-1")


def _pdf_bloque(pdf, texto, size=11, bold=False, color=(0, 0, 0)):
    # Fuerza que el cursor vuelva siempre al margen izquierdo antes y
    # despues de escribir, para evitar el error "Not enough horizontal
    # space" de fpdf2 cuando el cursor queda pegado al margen derecho.
    pdf.set_x(pdf.l_margin)
    pdf.set_font("Helvetica", "B" if bold else "", size)
    pdf.set_text_color(*color)
    pdf.multi_cell(
        pdf.w - pdf.l_margin - pdf.r_margin,
        size * 0.6,
        texto,
        new_x="LMARGIN",
        new_y="NEXT",
    )


def construir_pdf(titulo: str, secciones) -> bytes:
    """secciones: lista de tuplas (encabezado, cuerpo)."""
    pdf = FPDF()
    pdf.add_page()
    try:
        logo_bytes = base64.b64decode(LOGO_UNAE_B64)
        pdf.image(io.BytesIO(logo_bytes), x=pdf.l_margin, y=pdf.t_margin, h=16)
        # Deja espacio debajo del logo antes de seguir escribiendo texto
        pdf.set_xy(pdf.l_margin, pdf.t_margin + 20)
    except Exception:
        # Si el logo no se puede insertar por algún motivo, el PDF se genera igual sin él
        pass
    _pdf_bloque(pdf, _sanear(titulo), size=14, bold=True, color=(76, 1, 62))
    _pdf_bloque(
        pdf,
        datetime.now(ZoneInfo("America/Guayaquil")).strftime("Generado el %d/%m/%Y a las %H:%M"),
        size=9,
        color=(100, 100, 100),
    )
    pdf.ln(4)
    for encabezado, cuerpo in secciones:
        _pdf_bloque(pdf, _sanear(encabezado), size=11, bold=True, color=(18, 48, 92))
        _pdf_bloque(pdf, _sanear(cuerpo), size=11)
        pdf.ln(3)
    raw = pdf.output(dest="S")
    if isinstance(raw, (bytes, bytearray)):
        return bytes(raw)
    return raw.encode("latin-1")


def extraer_texto_archivo(archivo) -> str:
    """Extrae texto de un archivo subido por el estudiante (.txt, .docx o .pdf)."""
    nombre = archivo.name.lower()
    try:
        if nombre.endswith(".txt"):
            return archivo.read().decode("utf-8", errors="replace")
        elif nombre.endswith(".docx"):
            from docx import Document
            doc = Document(archivo)
            return "\n".join(p.text for p in doc.paragraphs)
        elif nombre.endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(archivo)
            return "\n".join((pagina.extract_text() or "") for pagina in reader.pages)
        else:
            return ""
    except Exception as e:
        st.error(f"No se pudo leer el archivo: {e}")
        return ""


tab1, tab2 = st.tabs(["📖 Consultar el manual", "✍️ Revisar mi texto académico"])

# ============================================================
# TAB 1: Preguntas y respuestas sobre el manual (RAG)
# ============================================================
with tab1:
    pregunta = st.text_area("Escriba su pregunta o consulta:", height=100, key="pregunta_manual")

    col1, col2 = st.columns([1, 5])
    with col1:
        enviar = st.button("Consultar", type="primary", key="btn_consultar")

    if enviar and pregunta.strip():
        with st.spinner("Consultando la base de conocimiento..."):
            try:
                resp = requests.post(
                    f"{CEDIA_BASE_URL}/rag",
                    headers={"x-api-key": API_KEY, "Content-Type": "application/json"},
                    json={
                        "question": pregunta,
                        "collection": COLLECTION,
                        "use_rerank": True,
                        "max_tokens": 1200,
                    },
                    timeout=60,
                )
                if resp.status_code == 200:
                    data = resp.json()
                    respuesta = data.get("answer", "(sin respuesta)")
                    st.session_state.historial.insert(0, (pregunta, respuesta))
                else:
                    st.error(f"Error {resp.status_code}: {resp.text}")
            except requests.exceptions.RequestException as e:
                st.error(f"No se pudo conectar con el servicio: {e}")

    st.divider()

    for i, (q, a) in enumerate(st.session_state.historial):
        st.markdown(
            f"""
            <div class="qa-card">
                <div class="qa-question">❓ {html.escape(q)}</div>
                <div class="qa-answer">{html.escape(a)}</div>
            </div>
            """,
            unsafe_allow_html=True,
        )
        pdf_bytes = construir_pdf(
            "Asistente de Escritura Academica - UNAE",
            [("Pregunta:", q), ("Respuesta:", a)],
        )
        st.download_button(
            label="📄 Descargar esta respuesta en PDF",
            data=pdf_bytes,
            file_name=f"respuesta_{i+1}.pdf",
            mime="application/pdf",
            key=f"pdf_{i}",
        )


# ============================================================
# MMCD — Parte 1 (trabajo salvavida)
# Configuración, registro de actividad y validación de salida
# ============================================================

# Dimensiones del MMCD. IMPORTANTE: las descripciones deben reemplazarse por
# las definiciones exactas de MMCD_normas_sistema_v1.json (fuente de verdad).
# Las de abajo son un resumen provisional tomado del plan de trabajo.
DIMENSIONES = {
    "EP": "Existencia de una tesis o postura clara del autor",
    "RCF": "Uso crítico de las fuentes (no solo reproducirlas)",
    "CCA": "Consideración de objeciones o contraargumentos",
    "TAS": "Avance o transformación del texto respecto de la versión anterior",
    "JD": "Justificación de las decisiones del autor",
}

SECCIONES = [
    "Planteamiento del problema",
    "Pregunta de investigación",
    "Objetivos",
    "Justificación",
    "Marco teórico",
    "Marco metodológico",
    "Análisis de datos",
    "Otro",
]

ETAPAS = ["Borrador", "Revisión", "Versión final"]

ENCABEZADO_FORMAL = "SEÑALAMIENTOS FORMALES"
ENCABEZADO_CRITICO = "PREGUNTAS DE MEDIACIÓN CRÍTICA"
MAX_CITA = 200
LIMITE_CARACTERES_REVISION = 50000

ARCHIVO_REGISTRO = Path(__file__).parent / "registro_actividad.csv"
COLUMNAS_REGISTRO = [
    "id_estudiante",          # hash SHA-256 truncado; nunca el dato original
    "fecha_hora",             # hora de Ecuador (America/Guayaquil)
    "tipo_documento",
    "seccion",
    "etapa",
    "criterios",
    "longitud_texto",         # caracteres efectivamente enviados al modelo
    "texto_recortado",
    "tiempo_respuesta_s",     # tiempo total, incluido el reintento si lo hubo
    "tiempo_primer_intento_s",
    "reintento",
    "formato_valido",
    "lineas_descartadas",
    "estado",                 # ok / error_http_XXX / error_conexion
]


def anonimizar(identificador: str) -> str:
    """Devuelve un identificador seudonimizado y estable para el mismo estudiante."""
    sal = st.secrets.get("SAL_ANONIMIZACION", "mmcd-unae-2026")
    base = f"{sal}|{identificador.strip().lower()}".encode("utf-8")
    return hashlib.sha256(base).hexdigest()[:12]


def _hoja_registro():
    """Devuelve la hoja de Google Sheets si está configurada en Secrets; si no, None."""
    if "gcp_service_account" not in st.secrets or "REGISTRO_SHEET_ID" not in st.secrets:
        return None
    try:
        import gspread
        cliente = gspread.service_account_from_dict(dict(st.secrets["gcp_service_account"]))
        hoja = cliente.open_by_key(st.secrets["REGISTRO_SHEET_ID"]).sheet1
        if not hoja.row_values(1):
            hoja.append_row(COLUMNAS_REGISTRO)
        return hoja
    except Exception as e:
        st.session_state["_error_hoja"] = str(e)
        return None


def registrar_actividad(fila: dict) -> None:
    """Guarda un renglón por consulta: en Google Sheets (persistente) si está
    configurado, y siempre en un CSV local como respaldo."""
    valores = [fila.get(c, "") for c in COLUMNAS_REGISTRO]
    hoja = _hoja_registro()
    if hoja is not None:
        try:
            hoja.append_row(valores, value_input_option="USER_ENTERED")
        except Exception as e:
            st.session_state["_error_hoja"] = str(e)
    try:
        nuevo = not ARCHIVO_REGISTRO.exists()
        with open(ARCHIVO_REGISTRO, "a", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            if nuevo:
                w.writerow(COLUMNAS_REGISTRO)
            w.writerow(valores)
    except Exception:
        pass


def _normalizar(linea: str) -> str:
    """Quita viñetas y marcas de Markdown al inicio y al final de una línea."""
    linea = linea.strip()
    linea = re.sub(r"^([-*•·]|\d+[.)])\s*", "", linea)
    return linea.strip().strip("*_ ").strip()


def _es_encabezado(linea: str, encabezado: str) -> bool:
    limpio = _normalizar(linea).lstrip("#").strip().rstrip(":").strip().upper()
    sin_tildes = lambda s: s.translate(str.maketrans("ÁÉÍÓÚÑ", "AEIOUN"))
    return sin_tildes(limpio) == sin_tildes(encabezado)


def separar_bloques(respuesta: str):
    """Divide la respuesta del modelo en (bloque_formal, lineas_criticas)."""
    formal, critico, actual = [], [], None
    for linea in respuesta.splitlines():
        if _es_encabezado(linea, ENCABEZADO_FORMAL):
            actual = formal
            continue
        if _es_encabezado(linea, ENCABEZADO_CRITICO):
            actual = critico
            continue
        if actual is not None and linea.strip():
            actual.append(_normalizar(linea))
    return "\n".join(formal).strip(), [l for l in critico if l]


def validar_criticas(lineas, criterios):
    """Una línea es válida si menciona una etiqueta seleccionada y termina en '?'."""
    patron = re.compile(r"\[(%s)\]" % "|".join(criterios))
    validas, descartadas = [], []
    for l in lineas:
        if patron.search(l) and l.rstrip(" »\"”)").endswith("?"):
            validas.append(l)
        else:
            descartadas.append(l)
    return validas, descartadas


def construir_instrucciones(tipo_doc, seccion, etapa, criterios, incluir_formal, texto):
    criterios_txt = "\n".join(f"   - [{c}] {DIMENSIONES[c]}" for c in criterios)
    if incluir_formal:
        bloque_formal = f"""{ENCABEZADO_FORMAL}:
- (un señalamiento por línea: cita el fragmento entre « » e indica la norma
  de ortografía, gramática, puntuación o APA que no cumple. No escribas la
  versión corregida.)"""
    else:
        bloque_formal = f"""{ENCABEZADO_FORMAL}:
- No solicitado."""

    return f"""Eres un mediador de escritura académica que aplica el Modelo de Mediación
Crítica Declarada (MMCD) de la Universidad Nacional de Educación (UNAE).
No eres corrector ni redactor: solo señalas y preguntas.

CONTEXTO DECLARADO POR EL ESTUDIANTE
- Tipo de documento: {tipo_doc}
- Sección del texto: {seccion}
- Etapa del proceso: {etapa}
- Criterios que el estudiante quiere priorizar:
{criterios_txt}

REGLAS OBLIGATORIAS
1. NUNCA reescribas, completes, parafrasees como propuesta ni resuelvas el
   texto o el argumento del estudiante. No entregues redacción alternativa,
   oraciones de ejemplo ni frases para copiar y pegar.
2. En el bloque {ENCABEZADO_CRITICO} no emitas juicios ni afirmaciones
   (por ejemplo, "esto debilita el argumento" o "parece apresurado").
   Toda oración de ese bloque es una pregunta y termina en "?".
3. Cada pregunta ocupa UNA sola línea con este formato exacto:
   [ETIQUETA] «cita textual del estudiante» — ¿pregunta?
   - ETIQUETA es una de las etiquetas priorizadas ({", ".join(criterios)}).
   - La cita se copia literalmente del texto del estudiante y tiene como
     máximo {MAX_CITA} caracteres.
4. Formula entre una y tres preguntas por cada criterio priorizado, adecuadas
   a la sección ({seccion}) y a la etapa ({etapa}). No uses otras etiquetas.
5. No mezcles capas: la gramática, la ortografía y el formato APA van solo
   en {ENCABEZADO_FORMAL}; la argumentación va solo en {ENCABEZADO_CRITICO}.
   Un mismo fragmento no recibe señalamiento formal y pregunta crítica por
   el mismo motivo.
6. Responde exactamente con estos dos bloques, en este orden, sin texto
   adicional antes, entre ni después:

{bloque_formal}
{ENCABEZADO_CRITICO}:
[ETIQUETA] «...» — ¿...?

Texto del estudiante:
\"\"\"
{texto}
\"\"\"
"""


def llamar_rag(pregunta: str, max_tokens: int = 2000, timeout: int = 240):
    """Devuelve (texto, estado, segundos)."""
    inicio = time.perf_counter()
    try:
        resp = requests.post(
            f"{CEDIA_BASE_URL}/rag",
            headers={"x-api-key": API_KEY, "Content-Type": "application/json"},
            json={
                "question": pregunta,
                "collection": COLLECTION,
                "use_rerank": True,
                "max_tokens": max_tokens,
            },
            timeout=timeout,
        )
        segundos = round(time.perf_counter() - inicio, 2)
        if resp.status_code != 200:
            return resp.text, f"error_http_{resp.status_code}", segundos
        data = resp.json()
        if isinstance(data, dict):
            texto = (
                data.get("answer") or data.get("response") or data.get("text")
                or data.get("content") or data.get("message") or data.get("output")
                or str(data)
            )
        else:
            texto = str(data)
        return texto, "ok", segundos
    except requests.exceptions.RequestException as e:
        return str(e), "error_conexion", round(time.perf_counter() - inicio, 2)


# ---------- PDF de retroalimentación (MMCD), ordenado y con nota de cierre ----------
COLOR_GUINDA = (76, 1, 62)
COLOR_GRIS = (95, 95, 95)
COLOR_FONDO = (246, 240, 244)

NOTA_CIERRE = (
    "El presente informe tiene carácter formativo y no constituye una calificación "
    "ni sustituye la revisión de su tutor o docente. Conforme al Modelo de Mediación "
    "Crítica Declarada (MMCD), el asistente no reescribe ni corrige su texto: señala "
    "aspectos a revisar y formula preguntas para orientar su reflexión. La decisión "
    "sobre qué cambios realizar, y su justificación, corresponde exclusivamente a "
    "usted como autor o autora del trabajo."
)


class PDFInforme(FPDF):
    """PDF con pie de página institucional y numeración en todas las páginas."""

    def footer(self):
        self.set_y(-15)
        self.set_draw_color(*COLOR_GUINDA)
        self.line(self.l_margin, self.get_y(), self.w - self.r_margin, self.get_y())
        self.ln(1.5)
        self.set_font("Helvetica", "", 8)
        self.set_text_color(*COLOR_GRIS)
        self.cell(0, 5, _sanear("Asistente de Escritura Académica · UNAE · Piloto RAG con CEDIA"), align="L")
        self.set_x(self.l_margin)
        self.cell(0, 5, f"Página {self.page_no()} de {{nb}}", align="R")


def _ancho_util(pdf):
    return pdf.w - pdf.l_margin - pdf.r_margin


def _titulo_seccion(pdf, numero, texto):
    if pdf.get_y() > pdf.h - 45:
        pdf.add_page()
    pdf.ln(3)
    pdf.set_x(pdf.l_margin)
    pdf.set_font("Helvetica", "B", 12)
    pdf.set_text_color(*COLOR_GUINDA)
    pdf.cell(0, 7, _sanear(f"{numero}. {texto}"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_draw_color(*COLOR_GUINDA)
    pdf.line(pdf.l_margin, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y())
    pdf.ln(3)


def _item_numerado(pdf, numero, texto, sangria=8):
    pdf.set_x(pdf.l_margin)
    pdf.set_font("Helvetica", "B", 10.5)
    pdf.set_text_color(*COLOR_GUINDA)
    pdf.cell(sangria, 5.6, f"{numero}.")
    pdf.set_font("Helvetica", "", 10.5)
    pdf.set_text_color(0, 0, 0)
    pdf.multi_cell(_ancho_util(pdf) - sangria, 5.6, _sanear(texto), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(1.5)


def _parsear_pregunta(linea: str):
    """Separa '[EP] «cita» — ¿pregunta?' en (etiqueta, cita, pregunta)."""
    m = re.match(r"^\s*\[(\w+)\]\s*«(.+?)»\s*[—–\-:]*\s*(.+)$", linea)
    if not m:
        m2 = re.match(r"^\s*\[(\w+)\]\s*(.+)$", linea)
        return (m2.group(1), "", m2.group(2).strip()) if m2 else ("", "", linea.strip())
    etiqueta, cita, pregunta = m.group(1), m.group(2).strip(), m.group(3).strip()
    # Primera letra en mayúscula después de "¿" (el modelo suele escribirla en minúscula)
    pregunta = re.sub(r"^(¿\s*)(\w)", lambda x: x.group(1) + x.group(2).upper(), pregunta)
    return etiqueta, cita, pregunta


def construir_pdf_retro(r: dict) -> bytes:
    pdf = PDFInforme()
    pdf.set_auto_page_break(auto=True, margin=22)
    pdf.add_page()
    ancho = _ancho_util(pdf)

    # Encabezado
    try:
        pdf.image(io.BytesIO(base64.b64decode(LOGO_UNAE_B64)), x=pdf.l_margin, y=pdf.t_margin, h=16)
        pdf.set_xy(pdf.l_margin, pdf.t_margin + 20)
    except Exception:
        pass
    pdf.set_font("Helvetica", "B", 15)
    pdf.set_text_color(*COLOR_GUINDA)
    pdf.cell(0, 8, _sanear("Informe de retroalimentación de escritura académica"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(*COLOR_GRIS)
    ahora = datetime.now(ZoneInfo("America/Guayaquil"))
    pdf.cell(0, 5, _sanear(ahora.strftime("Generado el %d/%m/%Y a las %H:%M (hora de Ecuador)")), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    # 1. Datos del texto revisado (tabla)
    _titulo_seccion(pdf, 1, "Datos del texto revisado")
    criterios_txt = "\n".join(f"{c} — {DIMENSIONES.get(c, '')}" for c in r["criterios"])
    filas = [
        ("Tipo de documento", r["tipo"]),
        ("Sección", r["seccion"]),
        ("Etapa del proceso", r["etapa"]),
        ("Criterios priorizados", criterios_txt),
    ]
    col1 = 45
    pdf.set_draw_color(220, 210, 218)
    for etiqueta, valor in filas:
        y0 = pdf.get_y()
        pdf.set_font("Helvetica", "", 10)
        alto = max(7, len(pdf.multi_cell(ancho - col1, 5.5, _sanear(valor), dry_run=True, output="LINES")) * 5.5 + 1.5)
        pdf.set_fill_color(*COLOR_FONDO)
        pdf.set_font("Helvetica", "B", 10)
        pdf.set_text_color(*COLOR_GUINDA)
        pdf.set_xy(pdf.l_margin, y0)
        pdf.multi_cell(col1, alto, _sanear(etiqueta), border=1, fill=True)
        pdf.set_xy(pdf.l_margin + col1, y0)
        pdf.set_font("Helvetica", "", 10)
        pdf.set_text_color(0, 0, 0)
        pdf.rect(pdf.l_margin + col1, y0, ancho - col1, alto)
        pdf.set_xy(pdf.l_margin + col1 + 1.5, y0 + 0.75)
        pdf.multi_cell(ancho - col1 - 3, 5.5, _sanear(valor))
        pdf.set_xy(pdf.l_margin, y0 + alto)

    # 2. Señalamientos formales
    _titulo_seccion(pdf, 2, "Señalamientos formales (ortografía, gramática y normas APA)")
    formales = [l.strip() for l in r["formal"].splitlines() if l.strip()]
    for i, linea in enumerate(formales, 1):
        _item_numerado(pdf, i, linea)

    # 3. Preguntas de mediación crítica, agrupadas por criterio
    _titulo_seccion(pdf, 3, "Preguntas de mediación crítica")
    if not r["criticas"]:
        _item_numerado(pdf, 1, "No se generaron preguntas válidas en esta ocasión.")
    grupos = {}
    for linea in r["criticas"]:
        etiqueta, cita, pregunta = _parsear_pregunta(linea)
        grupos.setdefault(etiqueta, []).append((cita, pregunta))
    n = 0
    for etiqueta, items in grupos.items():
        if pdf.get_y() > pdf.h - 50:
            pdf.add_page()
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "B", 10.5)
        pdf.set_text_color(*COLOR_GUINDA)
        nombre = DIMENSIONES.get(etiqueta, "")
        pdf.multi_cell(ancho, 6, _sanear(f"Criterio {etiqueta}" + (f" — {nombre}" if nombre else "")), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1)
        for cita, pregunta in items:
            n += 1
            if pdf.get_y() > pdf.h - 40:
                pdf.add_page()
            sangria = 8
            pdf.set_x(pdf.l_margin)
            pdf.set_font("Helvetica", "B", 10.5)
            pdf.cell(sangria, 5.6, f"{n}.")
            if cita:
                pdf.set_font("Helvetica", "I", 9.5)
                pdf.set_text_color(*COLOR_GRIS)
                pdf.multi_cell(ancho - sangria, 5.2, _sanear(f"Fragmento citado: «{cita}»"), new_x="LMARGIN", new_y="NEXT")
                pdf.set_x(pdf.l_margin + sangria)
            pdf.set_font("Helvetica", "", 10.5)
            pdf.set_text_color(0, 0, 0)
            pdf.multi_cell(ancho - sangria, 5.6, _sanear(pregunta), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2.5)

    # Nota de cierre (recuadro y firma se mantienen juntos en la misma página)
    pdf.set_font("Helvetica", "", 9.5)
    lineas = pdf.multi_cell(ancho - 8, 5, _sanear(NOTA_CIERRE), dry_run=True, output="LINES")
    alto = len(lineas) * 5 + 8
    if pdf.get_y() + 4 + 6 + alto + 5 + 16 > pdf.h - 22:
        pdf.add_page()
    pdf.ln(4)
    pdf.set_x(pdf.l_margin)
    pdf.set_font("Helvetica", "B", 10.5)
    pdf.set_text_color(*COLOR_GUINDA)
    pdf.cell(0, 6, "Nota de cierre", new_x="LMARGIN", new_y="NEXT")
    y0 = pdf.get_y()
    pdf.set_font("Helvetica", "", 9.5)
    pdf.set_fill_color(*COLOR_FONDO)
    pdf.set_draw_color(*COLOR_GUINDA)
    pdf.rect(pdf.l_margin, y0, ancho, alto, style="DF")
    pdf.set_xy(pdf.l_margin + 4, y0 + 4)
    pdf.set_text_color(40, 40, 40)
    pdf.multi_cell(ancho - 8, 5, _sanear(NOTA_CIERRE), align="J", new_x="LMARGIN", new_y="NEXT")
    pdf.set_y(y0 + alto + 5)
    pdf.set_font("Helvetica", "I", 9)
    pdf.set_text_color(*COLOR_GRIS)
    pdf.multi_cell(ancho, 5, _sanear(
        "Atentamente,\nAsistente de Escritura Académica\nUniversidad Nacional de Educación — UNAE"
    ), align="C", new_x="LMARGIN", new_y="NEXT")

    return bytes(pdf.output())


# ============================================================
# TAB 2: Retroalimentación sobre el texto del propio estudiante
# El programa NUNCA reescribe el texto; solo señala y pregunta (MMCD).
# ============================================================
with tab2:
    st.markdown(
        "Suba o pegue un **fragmento** de su ensayo, artículo científico o proyecto "
        "de titulación. El asistente le devuelve **señalamientos formales** y "
        "**preguntas de mediación crítica** — **no reescribe el texto por usted**; "
        "la decisión final sobre cómo mejorarlo es siempre suya."
    )

    id_estudiante_raw = st.text_input(
        "Correo institucional o código de estudiante *",
        key="id_estudiante",
        help="Se guarda seudonimizado (cifrado de un solo sentido). No se registra "
             "su nombre ni su correo.",
    )

    tipo_doc = st.selectbox(
        "Tipo de documento",
        ["Ensayo", "Artículo científico", "Proyecto de titulación (tesis)"],
        key="tipo_doc",
    )

    col_a, col_b = st.columns(2)
    with col_a:
        seccion = st.selectbox(
            "Sección del texto *", SECCIONES, index=None,
            placeholder="Seleccione la sección", key="seccion",
        )
    with col_b:
        etapa = st.selectbox(
            "Etapa del proceso *", ETAPAS, index=None,
            placeholder="Seleccione la etapa", key="etapa",
        )

    seccion_otro = ""
    if seccion == "Otro":
        seccion_otro = st.text_input("Especifique la sección", key="seccion_otro")

    criterios = st.multiselect(
        "Criterios que desea priorizar * (uno o varios)",
        list(DIMENSIONES.keys()),
        format_func=lambda c: f"{c} — {DIMENSIONES[c]}",
        key="criterios",
    )

    archivo = st.file_uploader(
        "Suba su archivo (.docx, .pdf o .txt) — opcional, también puede pegar el texto abajo",
        type=["docx", "pdf", "txt"],
        key="archivo_estudiante",
    )

    # Nota técnica: st.text_area ignora value= cuando ya existe su key en
    # session_state; por eso el texto extraído se escribe directamente ahí.
    if archivo is not None:
        identificador_archivo = f"{archivo.name}_{archivo.size}"
        if st.session_state.get("_ultimo_archivo_procesado") != identificador_archivo:
            texto_extraido = extraer_texto_archivo(archivo)
            st.session_state["_ultimo_archivo_procesado"] = identificador_archivo
            if texto_extraido:
                st.session_state["texto_estudiante"] = texto_extraido
                st.success(f"Se extrajeron {len(texto_extraido)} caracteres de '{archivo.name}'.")
            else:
                st.warning(
                    f"No se pudo extraer texto de '{archivo.name}'. Si es un PDF "
                    "escaneado (fotos de las páginas sin texto seleccionable), no "
                    "hay texto que leer — pruebe subiendo el archivo en Word "
                    "(.docx) o un PDF con texto real."
                )

    texto_estudiante = st.text_area(
        "O pegue aquí el texto:",
        height=280,
        key="texto_estudiante",
    )

    if len(texto_estudiante) > 30000:
        st.warning(
            f"El texto tiene {len(texto_estudiante):,} caracteres. Si supera "
            f"{LIMITE_CARACTERES_REVISION:,} caracteres, el sistema revisará solo "
            "la primera parte. Se recomienda enviar un fragmento por vez "
            "(por ejemplo, una sección)."
        )

    incluir_formal = st.checkbox(
        "Incluir señalamientos formales (ortografía, gramática y normas APA)",
        value=True,
        key="revisar_apa",
    )

    faltantes = []
    if not id_estudiante_raw.strip():
        faltantes.append("correo o código de estudiante")
    if not seccion:
        faltantes.append("sección")
    if not etapa:
        faltantes.append("etapa")
    if not criterios:
        faltantes.append("al menos un criterio")
    if not texto_estudiante.strip():
        faltantes.append("texto")

    st.info("⏳ Este análisis puede tardar hasta dos minutos; no cierre ni recargue la página.")

    if faltantes:
        st.caption("Para habilitar la revisión, complete: " + ", ".join(faltantes) + ".")

    revisar = st.button(
        "Revisar mi texto", type="primary", key="btn_revisar",
        disabled=bool(faltantes),
    )

    if revisar and not faltantes:
        seccion_final = f"Otro: {seccion_otro.strip()}" if seccion == "Otro" and seccion_otro.strip() else seccion

        texto_para_revisar = texto_estudiante
        fue_recortado = False
        if len(texto_para_revisar) > LIMITE_CARACTERES_REVISION:
            corte = texto_para_revisar.rfind(" ", 0, LIMITE_CARACTERES_REVISION)
            texto_para_revisar = texto_para_revisar[: corte if corte != -1 else LIMITE_CARACTERES_REVISION]
            fue_recortado = True
            st.warning(
                f"Se revisaron solo los primeros {len(texto_para_revisar):,} de "
                f"{len(texto_estudiante):,} caracteres. Envíe el resto por separado."
            )

        instrucciones = construir_instrucciones(
            tipo_doc, seccion_final, etapa, criterios, incluir_formal, texto_para_revisar
        )

        fila = {
            "id_estudiante": anonimizar(id_estudiante_raw),
            "fecha_hora": datetime.now(ZoneInfo("America/Guayaquil")).strftime("%Y-%m-%d %H:%M:%S"),
            "tipo_documento": tipo_doc,
            "seccion": seccion_final,
            "etapa": etapa,
            "criterios": "|".join(criterios),
            "longitud_texto": len(texto_para_revisar),
            "texto_recortado": "si" if fue_recortado else "no",
            "reintento": "no",
            "formato_valido": "",
            "lineas_descartadas": 0,
        }

        with st.spinner("Analizando el texto… puede tardar hasta dos minutos."):
            respuesta, estado, t1 = llamar_rag(instrucciones)
            tiempo_total = t1
            fila["tiempo_primer_intento_s"] = t1

            if estado == "ok":
                formal, lineas = separar_bloques(respuesta)
                validas, descartadas = validar_criticas(lineas, criterios)

                # Un reintento si el modelo no respetó el formato (G3/G8).
                if not lineas or descartadas:
                    fila["reintento"] = "si"
                    recordatorio = (
                        "ATENCIÓN: tu respuesta anterior no respetó el formato. En el "
                        f"bloque {ENCABEZADO_CRITICO} cada línea debe empezar con una "
                        "etiqueta entre corchetes, citar el texto entre « » y terminar "
                        "en '?'. No escribas afirmaciones ni reescrituras.\n\n"
                    )
                    r2, e2, t2 = llamar_rag(recordatorio + instrucciones)
                    tiempo_total = round(t1 + t2, 2)
                    if e2 == "ok":
                        f2, l2 = separar_bloques(r2)
                        v2, d2 = validar_criticas(l2, criterios)
                        if len(v2) >= len(validas):
                            formal, lineas, validas, descartadas = f2, l2, v2, d2

                # Lo que siga sin cumplir el formato se descarta antes de mostrarse.
                fila["formato_valido"] = "si" if (validas and not descartadas) else "no"
                fila["lineas_descartadas"] = len(descartadas)
                if not incluir_formal:
                    formal = "No solicitado."

                st.session_state.historial_retro.insert(0, {
                    "tipo": tipo_doc,
                    "seccion": seccion_final,
                    "etapa": etapa,
                    "criterios": criterios,
                    "formal": formal or "Sin señalamientos formales.",
                    "criticas": validas,
                })
            else:
                st.error(f"El servicio respondió con un error ({estado}): {respuesta[:500]}")

        fila["tiempo_respuesta_s"] = tiempo_total
        fila["estado"] = estado
        registrar_actividad(fila)

    st.divider()

    for i, r in enumerate(st.session_state.historial_retro):
        criticas_txt = "\n".join(r["criticas"]) if r["criticas"] else (
            "No se generaron preguntas válidas en esta ocasión. Intente nuevamente "
            "o envíe un fragmento más breve."
        )
        st.markdown(
            f"""
            <div class="qa-card">
                <div class="qa-question">📝 {html.escape(r['tipo'])} · {html.escape(r['seccion'])} · {html.escape(r['etapa'])} · {html.escape(', '.join(r['criterios']))}</div>
                <div class="qa-question" style="margin-top:10px;">{ENCABEZADO_FORMAL.capitalize()}</div>
                <div class="qa-answer">{html.escape(r['formal'])}</div>
                <div class="qa-question" style="margin-top:14px;">{ENCABEZADO_CRITICO.capitalize()}</div>
                <div class="qa-answer">{html.escape(criticas_txt)}</div>
            </div>
            """,
            unsafe_allow_html=True,
        )
        pdf_bytes = construir_pdf_retro(r)
        st.download_button(
            label="📄 Descargar esta retroalimentación en PDF",
            data=pdf_bytes,
            file_name=f"retroalimentacion_{i+1}.pdf",
            mime="application/pdf",
            key=f"pdf_retro_{i}",
        )

# ============================================================
# Panel de registro (solo coordinación, protegido con clave)
# ============================================================
if "ADMIN_PASSWORD" in st.secrets:
    with st.sidebar.expander("🔒 Registro de actividad (coordinación)"):
        clave = st.text_input("Clave", type="password", key="clave_admin")
        if clave and clave == st.secrets["ADMIN_PASSWORD"]:
            if "_error_hoja" in st.session_state:
                st.warning(f"Google Sheets no disponible: {st.session_state['_error_hoja']}")
            if ARCHIVO_REGISTRO.exists():
                datos = ARCHIVO_REGISTRO.read_bytes()
                n_consultas = max(datos.count(b"\n") - 1, 0)
                st.caption(f"{n_consultas} consultas en el CSV local.")
                st.download_button(
                    "Descargar registro_actividad.csv", datos,
                    file_name="registro_actividad.csv", mime="text/csv",
                )
            else:
                st.caption("Aún no hay consultas registradas en el CSV local.")
